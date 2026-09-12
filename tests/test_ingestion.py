"""Phase 2: document loaders and discovery."""

from __future__ import annotations

import docx
import pytest

from complaint_processor.ingestion import (
    DocumentLoadError,
    UnsupportedFormatError,
    discover_documents,
    is_supported,
    load_document,
    normalize_text,
)

from .pdf_fixture import write_pdf

SAMPLE = """Customer Complaint Record
Reference: CMP-1001

Customer Name: Anita Deshpande
Email: anita.deshpande@example.com
Phone: +91 98200 11223
Resolution Provided: Replacement unit shipped.
Status: Resolved
"""


@pytest.fixture
def txt_file(tmp_path):
    path = tmp_path / "complaint_001.txt"
    path.write_text(SAMPLE, encoding="utf-8")
    return path


@pytest.fixture
def pdf_file(tmp_path):
    return write_pdf(tmp_path / "complaint_002.pdf", SAMPLE)


@pytest.fixture
def docx_file(tmp_path):
    path = tmp_path / "complaint_003.docx"
    document = docx.Document()
    document.add_paragraph("Customer Complaint Record")
    table = document.add_table(rows=0, cols=2)
    for label, value in [
        ("Customer Name", "Ravi Menon"),
        ("Email", "ravi.menon@example.com"),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    document.add_paragraph("Resolution Provided: Awaiting engineer visit.")
    document.save(path)
    return path


class TestLoaders:
    def test_txt_roundtrip(self, txt_file):
        doc = load_document(txt_file)
        assert doc.file_type == "txt"
        assert doc.filename == "complaint_001.txt"
        assert doc.stem == "complaint_001"
        assert "anita.deshpande@example.com" in doc.text
        assert doc.char_count > 0

    def test_txt_falls_back_to_cp1252(self, tmp_path):
        path = tmp_path / "legacy.txt"
        path.write_bytes("Name: José\nIssue: it didn’t start.\n".encode("cp1252"))
        doc = load_document(path)
        assert "José" in doc.text
        assert "didn’t" in doc.text

    def test_pdf_text_is_extracted(self, pdf_file):
        doc = load_document(pdf_file)
        assert doc.file_type == "pdf"
        assert "anita.deshpande@example.com" in doc.text

    def test_docx_captures_table_cells(self, docx_file):
        """A paragraph-only reader would silently drop the customer details."""
        doc = load_document(docx_file)
        assert doc.file_type == "docx"
        assert "Ravi Menon" in doc.text
        assert "ravi.menon@example.com" in doc.text
        assert "Awaiting engineer visit" in doc.text


class TestFailures:
    def test_empty_file_raises(self, tmp_path):
        path = tmp_path / "blank.txt"
        path.write_text("   \n\n  \n", encoding="utf-8")
        with pytest.raises(DocumentLoadError, match="no text"):
            load_document(path)

    def test_unsupported_extension_raises(self, tmp_path):
        path = tmp_path / "notes.md"
        path.write_text("not a complaint", encoding="utf-8")
        with pytest.raises(UnsupportedFormatError, match="No loader"):
            load_document(path)

    def test_corrupt_pdf_raises(self, tmp_path):
        path = tmp_path / "broken.pdf"
        path.write_bytes(b"this is not a pdf at all")
        with pytest.raises(DocumentLoadError):
            load_document(path)

    def test_missing_folder_raises(self, tmp_path):
        with pytest.raises(DocumentLoadError, match="not found"):
            discover_documents(tmp_path / "does_not_exist")


class TestDiscovery:
    def test_finds_supported_files_in_sorted_order(self, txt_file, pdf_file, docx_file):
        found = discover_documents(txt_file.parent)
        assert [p.name for p in found] == [
            "complaint_001.txt",
            "complaint_002.pdf",
            "complaint_003.docx",
        ]

    def test_skips_unsupported_and_hidden_files(self, tmp_path, txt_file):
        (tmp_path / "README.md").write_text("ignore me", encoding="utf-8")
        (tmp_path / ".gitkeep").write_text("", encoding="utf-8")
        (tmp_path / "subfolder").mkdir()

        found = discover_documents(tmp_path)
        assert [p.name for p in found] == ["complaint_001.txt"]

    def test_empty_folder_returns_empty_list(self, tmp_path):
        assert discover_documents(tmp_path) == []


class TestNormalization:
    def test_collapses_excess_blank_lines(self):
        assert normalize_text("a\n\n\n\n\nb") == "a\n\nb"

    def test_strips_trailing_spaces_and_normalizes_newlines(self):
        assert normalize_text("a   \r\nb\t\r\n") == "a\nb"

    def test_is_supported_ignores_case(self, tmp_path):
        assert is_supported(tmp_path / "X.PDF")
        assert not is_supported(tmp_path / "X.rtf")
