"""Generate the sample complaint documents in data/.

Run as a module so the package imports resolve:

    python -m tests.make_sample_data

The five cases are chosen to exercise different paths rather than to look
similar: a resolved case, an open escalation, a record missing contact details,
an enquiry that is not a complaint, and a details table in Word format.
"""

from __future__ import annotations

import docx

from complaint_processor.config import settings

from .pdf_fixture import write_pdf

COMPLAINT_001 = """CUSTOMER COMPLAINT RECORD
Reference: CMP-2024-001
Date Received: 04 March 2024

Customer Name: Anita Deshpande
Email: anita.deshpande@example.com
Phone: +91 98200 11223

Product / Service: Model X400 Water Purifier
Purchase Date: 18 February 2024

Complaint Description:
The purifier stopped dispensing water four days after installation. The
display showed a filter error even though the unit was brand new. The
customer called the helpline twice on 24 and 25 February.

Resolution Provided:
A replacement unit was shipped on 08 March and installed by an engineer on
12 March. The faulty unit was collected at the same visit. The customer
confirmed the replacement is working correctly.

Escalation: Not required.
Supporting Information: Copy of the original invoice attached.
Case Status: Closed - Resolved
"""

COMPLAINT_002 = """URGENT - THIRD FOLLOW UP
Reference: CMP-2024-002

Customer: Mohammed Faruqui
Contact Number: +91 90040 77321
Email: m.faruqui@example.com

Product: Model W900 Washing Machine

I am writing for the third time about the washing machine delivered on
29 January. It has leaked from the base since the day it was installed.
I have called your helpline four times. On the last call the agent told me
an engineer would attend within 48 hours. That was eleven days ago and
nobody has visited.

The water has now damaged the flooring in my utility room. I have attached
photographs of the damage and a copy of the delivery note.

I want a manager to call me today. If I do not hear back I will take this to
the consumer court.

Status: Open, no engineer visit completed.
"""

COMPLAINT_003 = """Customer Service Enquiry
Reference: ENQ-2024-013

From: Priya Raghavan
Email: priya.raghavan@example.com

Hello,

I am considering the Model X400 water purifier for a household of four
people. Before I order, could you tell me:

1. How often does the filter cartridge need replacing?
2. Is professional installation included in the listed price?
3. What does the warranty cover in the first year?

I am not reporting any problem, I would just like the details before
deciding.

Many thanks,
Priya
"""

COMPLAINT_004 = """SUPPORT TICKET 88213
Channel: Email
Received: 11 March 2024

Contact address on file: r.menon2291@example.com

Ticket body:
Bought an air purifier from your store last month. The filter replacement
indicator stays red even after fitting a new filter, and the reset button
does nothing. I have tried unplugging the unit overnight as the manual
suggests.

No name or phone number was supplied with this ticket and no attachments
were included.

Agent notes: No action taken yet, awaiting triage.
"""


def build() -> list[str]:
    data_dir = settings.data_dir
    data_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []

    # A plain text record and a PDF, the two most common export formats.
    (data_dir / "complaint_001.txt").write_text(COMPLAINT_001, encoding="utf-8")
    written.append("complaint_001.txt")

    write_pdf(data_dir / "complaint_002.pdf", COMPLAINT_002)
    written.append("complaint_002.pdf")

    (data_dir / "complaint_003.txt").write_text(COMPLAINT_003, encoding="utf-8")
    written.append("complaint_003.txt")

    write_pdf(data_dir / "complaint_004.pdf", COMPLAINT_004)
    written.append("complaint_004.pdf")

    # Word format, with the customer details in a table. A paragraph-only
    # reader would lose the name, email and phone entirely.
    document = docx.Document()
    document.add_heading("Customer Complaint Record", level=1)
    document.add_paragraph("Reference: CMP-2024-005")
    document.add_paragraph("Date Received: 09 March 2024")

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in [
        ("Customer Name", "Sunita Kulkarni"),
        ("Email", "sunita.kulkarni@example.com"),
        ("Phone", "+91 99100 55447"),
        ("Product", "Model B200 Air Purifier"),
        ("Order Reference", "ORD-55120"),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_paragraph("")
    document.add_paragraph("Complaint Description:")
    document.add_paragraph(
        "The air purifier arrived with a cracked front panel. The unit powers "
        "on and appears to function, but the damage is visible from the front "
        "of the room and the customer does not want to keep it in this state."
    )
    document.add_paragraph("Resolution Provided:")
    document.add_paragraph(
        "A replacement front panel has been ordered from the supplier. The "
        "customer was told it would arrive within two weeks and that an "
        "engineer would fit it. The part has not yet arrived."
    )
    document.add_paragraph("Escalation: Not required at this stage.")
    document.add_paragraph("Supporting Information: Photograph of the crack attached.")
    document.add_paragraph("Case Status: In progress, awaiting part.")
    document.save(data_dir / "complaint_005.docx")
    written.append("complaint_005.docx")

    return written


if __name__ == "__main__":
    for name in build():
        print("wrote", name)
